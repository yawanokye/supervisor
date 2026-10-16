import asyncio
import copy
import hashlib
import json
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.requests import Request

from app import background_review, main, openai_batch, worker
from app.ai_config import HybridAIConfig
from app.ai_providers import AIProviderError, OpenAIProvider
from app.ai_schemas import AIUsageRecord, AcademicReviewBatch
from app.database import Base, ReviewRecord
from app.model_router import CostAwareAIProvider
from app.processing_options import processing_config, processing_options
from app.review_timing import estimate_review

NOW = datetime.now(timezone.utc)


@pytest.fixture
def sessions(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    for module in (background_review, main, openai_batch, worker):
        monkeypatch.setattr(module, "SessionLocal", factory)
    yield factory
    engine.dispose()


def record(sessions, **extra):
    values = dict(job_id="a" * 32, lecturer_id=1, filename="thesis.docx", academic_level="PhD", research_approach="mixed",
        review_scope="full_thesis", review_depth="standard", submission_stage="initial", status="processing", progress=98,
        review_delivery="automatic", processing_mode="standard", guided_mode=True, guided_total_units=3,
        guided_current_index=0, guided_units_json="[1,2,3]", guided_review_ids_json="{}", lease_owner="worker-one",
        lease_expires_at=NOW + timedelta(minutes=5), recoverable=True, payload_available=True, estimated_pages=60, created_at=NOW)
    values.update(extra)
    with sessions() as db:
        row = ReviewRecord(**values)
        db.add(row); db.commit()
        return row


def finish():
    return background_review.finish_chapter_turn("a" * 32, worker_id="worker-one", current_index=0, units=[1,2,3],
        review_id="first-report", usage_tokens=1250, checkpoint_count=4, summary={"overall_score": 75})


def test_chapter_transition_durable_and_not_counted_twice(sessions):
    record(sessions)
    assert finish()["status"] == "queued"
    assert finish() is None
    with sessions() as db:
        row = db.query(ReviewRecord).first()
        assert (row.guided_current_index, row.guided_tokens_used, row.lease_owner) == (1, 1250, None)
        assert json.loads(row.guided_review_ids_json) == {"1": "first-report"}
        result = background_review.public_job_progress({}, row)
        assert result["progress"] == 34
        assert result["chapter_progress"] == 2


def test_guided_choice_preserves_manual_continue(sessions):
    record(sessions, review_delivery="guided")
    assert finish()["status"] == "awaiting_continue"
    with sessions() as db:
        assert db.query(ReviewRecord).first().guided_current_index == 0


@pytest.mark.parametrize("changes", [{"status":"stopped"}, {"lease_owner":"different-worker"}, {"guided_current_index":1}])
def test_late_chapter_cannot_overwrite_stop_or_other_lease(sessions, changes):
    record(sessions, **changes)
    assert finish() is None
    with sessions() as db:
        assert db.query(ReviewRecord).first().guided_tokens_used == 0


def test_worker_continues_new_chapter_without_spending_recovery_budget(sessions, monkeypatch):
    record(sessions)
    finish()
    monkeypatch.setattr(worker, "payload_available", lambda job: True)
    assert worker._next_job_id() == ("a" * 32, False)


def test_batch_limit_keeps_capacity_for_routine_reviews(sessions, monkeypatch):
    record(sessions, processing_mode="batch")
    record(sessions, job_id="b"*32, status="queued", processing_mode="batch")
    record(sessions, job_id="c"*32, status="queued", processing_mode="flex")
    monkeypatch.setenv("VPROF_BATCH_WORKER_LIMIT", "1")
    monkeypatch.setattr(worker, "payload_available", lambda job: True)
    assert worker._next_job_id()[0] == "c" * 32


def test_job_can_be_reopened_without_browser_cache_and_remains_authorised(sessions, monkeypatch):
    row = record(sessions)
    main.JOB_CACHE.pop(row.job_id, None)
    monkeypatch.setattr(main, "_normalise_stalled_recovery_record", lambda db, rec: rec)
    monkeypatch.setattr(main, "_current_user", lambda req, db: SimpleNamespace(id=1, role="lecturer"))
    request = Request({"type":"http", "headers":[]})
    with sessions() as db:
        result = asyncio.run(main.get_review_job(row.job_id, request, db))
        assert result["progress"] < 100 and result["review_delivery"] == "automatic"
    monkeypatch.setattr(main, "_current_user", lambda req, db: SimpleNamespace(id=2, role="lecturer"))
    with sessions() as db, pytest.raises(main.HTTPException) as error:
        asyncio.run(main.get_review_job(row.job_id, request, db))
    assert error.value.status_code == 404


def test_estimate_accounts_for_queue_and_never_promises_fixed_twelve_hours(sessions):
    record(sessions, guided_mode=False, progress=0, started_at=None)
    row = record(sessions, job_id="b"*32, status="queued", guided_mode=False, progress=0, created_at=NOW + timedelta(seconds=1))
    with sessions() as db:
        estimate = estimate_review(db.get(ReviewRecord, row.id), db, worker_slots=1, now=NOW)
    assert estimate["estimated_queue_seconds"] >= 3600
    assert estimate["estimated_seconds_remaining"] > estimate["estimated_queue_seconds"]
    assert estimate["estimate_is_approximate"] and estimate["check_back_at"]


def test_batch_estimate_uses_actual_window(sessions):
    row = record(sessions, processing_mode="batch", batch_deadline_at=NOW + timedelta(hours=15), progress=35)
    result = estimate_review(row, now=NOW)
    assert result["estimated_seconds_remaining"] >= 15*3600
    assert "beyond" in result["estimate_note"]
    assert result["batch_window_ends_at"] == row.batch_deadline_at.isoformat()


@pytest.mark.parametrize("status", ["completed", "failed", "stopped", "awaiting_continue"])
def test_terminal_jobs_have_no_false_ready_time(sessions, status):
    result = estimate_review(record(sessions, status=status), now=NOW)
    assert result["estimated_completion_at"] is None


def test_economy_config_and_actual_tier_costs(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    config = processing_config(HybridAIConfig.from_env(), "flex")
    assert config.openai_fast_model == "gpt-6-luna" and config.openai_chapter_model == "gpt-6.1-sol"
    router = CostAwareAIProvider(config)
    base = dict(provider="openai", model="gpt-6.1-sol", purpose="test", input_tokens=1000, output_tokens=1000)
    standard = router._price_usage(AIUsageRecord(**base)).estimated_cost_usd
    assert router._price_usage(AIUsageRecord(**base, service_tier="flex")).estimated_cost_usd == pytest.approx(standard/2)
    assert router._price_usage(AIUsageRecord(**base, service_tier="batch")).estimated_cost_usd == pytest.approx(standard/2)


@pytest.mark.parametrize("failure,allow,expected", [("HTTP 429 resource_unavailable", True, 2), ("HTTP 429 resource_unavailable", False, 1), ("HTTP 429 quota_exceeded", True, 1), ("HTTP 408 request timeout", True, 1)])
def test_flex_fallback_only_for_rejected_capacity(monkeypatch, failure, allow, expected):
    from dataclasses import replace
    import app.ai_providers as providers
    calls = []
    async def fake_post(**kwargs):
        calls.append(kwargs)
        if len(calls) == 1:
            raise AIProviderError(failure)
        return {"status":"completed", "output_text":'{"reviews":[]}', "usage":{"input_tokens":10,"output_tokens":5}}, "request"
    monkeypatch.setattr(providers, "_post_json_with_retry", fake_post)
    config = replace(HybridAIConfig.from_env(), structured_output_retries=0, max_retries=0)
    async def run():
        token = processing_options.set({"mode":"flex", "allow_standard_fallback":allow})
        try:
            return await OpenAIProvider(config).complete_json(model="gpt-6.1-sol", system_prompt="rubric", user_prompt="chapter", schema_model=AcademicReviewBatch, purpose="test")
        finally:
            processing_options.reset(token)
    if expected == 2:
        result = asyncio.run(run())
        assert result.usage.standard_fallback_used
        assert "service_tier" not in calls[1]["payload"]
    else:
        with pytest.raises(AIProviderError): asyncio.run(run())
    assert len(calls) == expected
    assert calls[0]["payload"]["service_tier"] == "flex"
    assert calls[0]["timeout_seconds"] >= 900 and not calls[0]["payload"].get("background")


class MemoryCheckpoints:
    def __init__(self): self.saved = {}
    def load(self, key, **kwargs): return copy.deepcopy(self.saved.get(key))
    def save(self, key, value, **kwargs): self.saved[key] = copy.deepcopy(value)


def packet(text): return {"model":"gpt-6.1-sol", "input":[{"role":"system","content":"grading rubric"},{"role":"user","content":text}], "text":{"format":{"type":"json_schema"}}}


def test_batch_jsonl_groups_requests_and_reuses_saved_unordered_results(sessions):
    checkpoints = MemoryCheckpoints(); calls=[]; jsonl=[]
    async def http(method, path, **kwargs):
        calls.append((method,path))
        if path == "/files":
            jsonl.extend(json.loads(line) for line in kwargs["files"]["file"][1].decode().splitlines())
            return {"id":"file-in"}
        if path == "/batches": return {"id":"batch-one"}
        if path == "/batches/batch-one": return {"status":"completed", "output_file_id":"file-out"}
        if path == "/files/file-out/content": return "\n".join(json.dumps({"custom_id":p["custom_id"],"response":{"status_code":200,"body":{"status":"completed","value":p["body"]["input"][1]["content"]}}}) for p in reversed(jsonl))
        raise AssertionError(path)
    async def run():
        session = openai_batch.BatchSession(HybridAIConfig.from_env(), checkpoints, "a"*32, flush_seconds=0)
        session.http = http
        result = await asyncio.gather(session.request(packet("one")), session.request(packet("two")))
        await session.close()
        restored = openai_batch.BatchSession(HybridAIConfig.from_env(), checkpoints, "a"*32, flush_seconds=0)
        restored.http = http
        again = await restored.request(packet("one"))
        await restored.close()
        assert again[0]["value"] == "one"
        return result
    results = asyncio.run(run())
    assert [r[0]["value"] for r in results] == ["one", "two"]
    assert calls.count(("POST","/batches")) == 1
    assert len(jsonl) == 2 and len({p["custom_id"] for p in jsonl}) == 2
    assert all(p["url"] == "/v1/responses" and p["body"]["input"][0]["content"] == "grading rubric" for p in jsonl)
    assert results[0][0]["_vprof_service_tier"] == "batch"


def test_uncertain_batch_acknowledgement_cannot_duplicate_submission():
    checkpoints=MemoryCheckpoints(); body=packet("one"); digest=hashlib.sha256(json.dumps(body, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    value=(json.dumps({"custom_id":digest,"method":"POST","url":"/v1/responses","body":body},ensure_ascii=False)+"\n").encode()
    key="batch-group-"+hashlib.sha256(value).hexdigest()[:32]
    checkpoints.save(key, {"submission_started":True,"input_file_id":"file-in"})
    calls=[]
    async def http(method,path,**kwargs): calls.append(method); return {"data":[],"has_more":False}
    async def run():
        session=openai_batch.BatchSession(HybridAIConfig.from_env(),checkpoints,"job");session.http=http
        with pytest.raises(openai_batch.BatchTransportError,match="not submitted twice"):
            await session.submit({digest:body})
    asyncio.run(run())
    assert calls == ["GET"]


def test_retry_failed_batch_packet_changes_group_and_keeps_success(sessions):
    checkpoints=MemoryCheckpoints(); body=packet("retry")
    digest=hashlib.sha256(json.dumps(body,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    checkpoints.save("batch-request-"+digest[:32],{"batch_id":"old"})
    checkpoints.save("batch-output-old",{"status":"expired","results":{}})
    calls=[]
    async def http(method,path,**kwargs):
        calls.append((method,path))
        if path=="/files":return {"id":"new-file"}
        if path=="/batches":return {"id":"new-batch"}
        if path=="/batches/new-batch":return {"status":"completed","output_file_id":"new-out"}
        if path=="/files/new-out/content":return json.dumps({"custom_id":digest,"response":{"status_code":200,"body":{"output_text":"recovered"}}})
        raise AssertionError(path)
    async def run():
        session=openai_batch.BatchSession(HybridAIConfig.from_env(),checkpoints,"job",flush_seconds=0);session.http=http
        result=await session.request(body);await session.close();return result
    assert asyncio.run(run())[0]["output_text"]=="recovered"
    assert checkpoints.load("batch-request-"+digest[:32])["batch_id"]=="new-batch"


@pytest.mark.parametrize("mode,delivery,guided", [("flex","automatic",True),("standard","guided",True),("batch","guided",False)])
def test_upload_persists_options_before_acknowledging_background_job(sessions, monkeypatch, mode, delivery, guided):
    import io
    from docx import Document
    from fastapi.testclient import TestClient
    document=Document()
    for title in ("CHAPTER ONE", "CHAPTER TWO"):
        document.add_heading(title, level=1);document.add_paragraph("The chapter discusses digital finance and small business performance.")
    stream=io.BytesIO();document.save(stream)
    events=[]; saved={}
    def save(job,payload):
        events.append("persist");saved.update(copy.deepcopy(payload));return {"document_hash":"d"*64,"payload_hash":"p"*64}
    def schedule(*args,**kwargs): events.append("schedule");return False
    monkeypatch.setenv("OPENAI_API_KEY","test")
    monkeypatch.setattr(main,"_verify_csrf",lambda *args:None)
    monkeypatch.setattr(main,"_current_user",lambda *args:SimpleNamespace(id=1,role="lecturer",must_change_password=False,full_name="Reviewer",department="Research"))
    monkeypatch.setattr(main,"save_job_payload",save)
    monkeypatch.setattr(main,"_schedule_review_job",schedule)
    monkeypatch.setattr(main,"reserve_review_tokens",lambda *args,**kwargs:False)
    def database():
        with sessions() as db: yield db
    main.app.dependency_overrides[main.get_db]=database
    try:
        client=TestClient(main.app)
        response=client.post("/api/review", data={"academic_level":"PhD","research_approach":"mixed","csrf_token":"test","selected_chapter":"1","processing_mode":mode,"review_delivery":delivery,"related_work_enabled":"true","related_work_sources":["openalex","crossref"]},files={"file":("thesis.docx",stream.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert response.status_code == 202, response.text
        result=response.json()
        assert events == ["persist","schedule"]
        assert result["background_processing"] and result["processing_mode"] == mode
        assert saved["guided_mode"] == guided
        assert saved["review_delivery"] == ("automatic" if mode=="batch" else delivery)
        assert saved["related_work_sources"] == ["openalex","crossref"]
        with sessions() as db:
            row=db.query(ReviewRecord).first()
            assert row.processing_mode == mode and row.review_delivery == saved["review_delivery"]
        main.JOB_CACHE.pop(result["job_id"],None)
    finally:
        main.app.dependency_overrides.pop(main.get_db,None)


def test_primary_provider_uses_batch_session_and_retains_actual_usage(monkeypatch):
    import app.ai_providers as providers
    class Transport:
        async def request(self, body):
            assert body["input"][0]["content"] == "rubric"
            return {"status":"completed","output_text":'{"reviews":[]}',"_vprof_service_tier":"batch","usage":{"input_tokens":200,"output_tokens":20}}, "batch:packet"
    async def unexpected(**kwargs): raise AssertionError("A primary Batch packet used the live endpoint.")
    monkeypatch.setattr(providers,"_post_json_with_retry",unexpected)
    async def run():
        token=processing_options.set({"mode":"batch","batch_session":Transport()})
        try:
            return await OpenAIProvider(HybridAIConfig.from_env()).complete_json(model="gpt-6.1-sol",system_prompt="rubric",user_prompt="chapter",schema_model=AcademicReviewBatch,purpose="primary")
        finally: processing_options.reset(token)
    result=asyncio.run(run())
    assert result.usage.service_tier == "batch" and result.usage.input_tokens == 200


def test_export_failure_recovers_saved_review_without_any_academic_call(sessions, monkeypatch):
    import io
    from docx import Document
    from app.annotated_exporter import build_annotated_docx
    source=Document();source.add_heading("CHAPTER ONE",1);source.add_paragraph("The study examines digital finance in Ghana.")
    stream=io.BytesIO();source.save(stream)
    row=dict(finding_id="missing",finding_number=1,status="does_not_meet_requirement",severity="major",category="chapter_structure",section="Definition of Terms",missing_section_label="Definition of Terms",section_status="missing",item="Definition of Terms section is missing",assessment="Definitions are absent.",required_action="Define the study constructs.",annotation_eligible=True,evidence=[])
    review={"review_id":"saved-final-review","summary":{"filename":"thesis.docx","annotated_document_available":True,"overall_score":60},"canonical_findings":[row]}
    record(sessions,status="queued",guided_mode=False,processing_mode="batch",lease_owner=None,lease_expires_at=None)
    manager=MemoryCheckpoints();manager.completed_count=lambda:len(manager.saved)
    manager.mark_failed=lambda *args,**kwargs:None
    manager.save("pipeline-final",{"review":review,"usage_snapshot":{}})
    monkeypatch.setattr(main,"CheckpointManager",lambda *args:manager)
    monkeypatch.setattr(main,"payload_available",lambda *args:True)
    def forbidden(*args,**kwargs): raise AssertionError("Export recovery repeated academic work.")
    for name in ("analyse","enrich_review_with_academic_ai","run_final_consistency_audit","polish_selected_comments"):
        monkeypatch.setattr(main,name,forbidden)
    for name in ("save_annotated","save_inline_annotated","save_clean_revised","save_review_json"):
        monkeypatch.setattr(main,name,lambda *args:None)
    attempts=[]
    def native(*args):
        attempts.append(1)
        if len(attempts)==1: raise RuntimeError("Synthetic export interruption")
        return build_annotated_docx(*args)
    monkeypatch.setattr(main,"build_annotated_docx",native)
    def failed(job_id,**kwargs):
        with sessions() as db:
            current=db.query(ReviewRecord).first();current.status="failed";current.lease_owner=None;current.lease_expires_at=None;db.commit()
        return False
    monkeypatch.setattr(main,"_queue_automatic_retry",failed)
    payload=dict(data=stream.getvalue(),filename="thesis.docx",document_hash="d"*64,payload_hash="p"*64,workflow_type="supervisory_review",processing_mode="batch",review_depth="standard",review_delivery="automatic")
    asyncio.run(main._run_review_job("a"*32,payload,resumed=True))
    assert manager.load("delivery-final")["review"]["review_id"]=="saved-final-review"
    assert manager.load("document-export-diagnostic")["step"]=="native Word comments"
    with sessions() as db:
        current=db.query(ReviewRecord).first();current.status="queued";db.commit()
    asyncio.run(main._run_review_job("a"*32,payload,resumed=True))
    with sessions() as db:
        assert db.query(ReviewRecord).first().status=="completed"
    assert len(attempts)==2
    main.JOB_CACHE.pop("a"*32,None)
    main.REVIEW_CACHE.pop("saved-final-review",None)


def test_stopped_job_cannot_be_claimed_by_a_late_scheduled_task(sessions):
    record(sessions,status="stopped",lease_owner=None,lease_expires_at=None)
    assert main._claim_job("a"*32,resumed=False) is False


def test_explicit_stop_requests_cancellation_of_active_batches(sessions, monkeypatch):
    from dataclasses import replace
    from app.database import ReviewCheckpoint
    import app.checkpointing as checkpointing
    manager=MemoryCheckpoints()
    manager.save("batch-group-a",{"batch_id":"active"})
    manager.save("batch-group-b",{"batch_id":"finished"})
    with sessions() as db:
        for name in ("batch-group-a","batch-group-b"):
            db.add(ReviewCheckpoint(job_id="job",stage_key=name,status="completed"))
        db.commit()
    monkeypatch.setattr(checkpointing,"CheckpointManager",lambda *args:manager)
    calls=[]
    async def http(self,method,path,**kwargs):
        calls.append((method,path));return {"status":"completed" if path.endswith("finished") else "in_progress"}
    monkeypatch.setattr(openai_batch.BatchSession,"http",http)
    config=replace(HybridAIConfig.from_env(),openai_api_key="test",enabled=True)
    result=asyncio.run(openai_batch.cancel_batches("job",config))
    assert result=={"cancelled":1,"failed":0}
    assert ("POST","/batches/active/cancel") in calls
    assert ("POST","/batches/finished/cancel") not in calls
