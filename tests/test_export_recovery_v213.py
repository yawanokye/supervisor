import copy
import io

import pytest
from docx import Document

from app.annotated_exporter import build_annotated_docx, native_annotation_audit
from app.inline_annotated_exporter import build_inline_annotated_docx, inline_annotation_audit


def missing_section():
    return dict(finding_id="M1", finding_number=1, status="does_not_meet_requirement", severity="major", category="chapter_structure",
        section="Definition of Terms", missing_section_label="Definition of Terms", section_status="missing", section_contract_verified=True,
        item="Definition of Terms section is missing", assessment="The required definitions are absent.",
        required_action="Add definitions of the study constructs.", annotation_eligible=True, evidence=[])


@pytest.mark.parametrize("grouped", ["true", "false"])
def test_absent_section_has_explicit_note_and_valid_native_and_inline_bundle(monkeypatch, grouped):
    monkeypatch.setenv("VPROF_NATIVE_COMMENT_STYLE", "grouped" if grouped == "true" else "individual")
    document=Document();document.add_heading("CHAPTER ONE",1)
    original="The study examines digital finance and small business performance in Ghana."
    document.add_paragraph(original);stream=io.BytesIO();document.save(stream)
    review={"summary":{},"canonical_findings":[missing_section()]}
    native=build_annotated_docx(stream.getvalue(),copy.deepcopy(review),"Reviewer")
    assert native_annotation_audit(native,review,comment_author="Reviewer")["passed"]
    restored=Document(io.BytesIO(native))
    assert restored.paragraphs[1].text == original
    assert any("Reviewer notes for missing sections" in p.text for p in restored.paragraphs)
    assert any("definitions" in c.text.lower() for c in restored.comments)
    inline=build_inline_annotated_docx(stream.getvalue(),copy.deepcopy(review),"Reviewer")
    assert inline_annotation_audit(inline,review)["passed"]
