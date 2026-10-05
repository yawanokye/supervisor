import copy
import io
import zipfile
import pytest

from docx import Document
from lxml import etree

from app.annotated_exporter import build_annotated_docx, _source_locator_map
from app.inline_annotated_exporter import build_inline_annotated_docx
from app.annotation_placement import resolve_verified_location
from app.document_parser import parse_document


def source(doc):
    buffer=io.BytesIO();doc.save(buffer);return buffer.getvalue()


def finding(evidence,quote,number=1):
    return {'finding_id':f'location-{number}','finding_number':number,'status':'partly_meets_requirement',
        'severity':'major','confidence':.9,'category':'methodological_rigour','item':'Verify the reported sampling claim',
        'assessment':'The sampling claim needs verification against the procedure used.',
        'required_action':'Correct the sampling description using the actual procedure.',
        'section':evidence.get('heading') or 'Research Methods','chapter_number':evidence.get('chapter_number'),
        'problematic_quote':quote,'evidence':[copy.deepcopy(evidence)],'annotation_eligible':True}


def ranges(data):
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    result={}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root=etree.fromstring(z.read('word/document.xml'))
    for paragraph in root.findall('.//w:p',ns):
        active=set()
        for node in paragraph.iter():
            if node.tag.endswith('}commentRangeStart'):
                active.add(node.get('{'+ns['w']+'}id'))
            elif node.tag.endswith('}commentRangeEnd'):
                active.discard(node.get('{'+ns['w']+'}id'))
            elif node.tag.endswith('}t'):
                for key in active:result[key]=result.get(key,'')+(node.text or '')
    return list(result.values())


def test_stale_ordinal_is_relocated_by_exact_evidence_not_nearby_words():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Research Design',2)
    doc.add_paragraph('This unrelated paragraph describes the study setting.')
    target='The sampling procedure selected only volunteers.'
    doc.add_paragraph(target)
    data=source(doc);rows=parse_document(data,'chapter.docx')
    evidence=next(r for r in rows if r['text']==target);evidence['paragraph']-=1
    review={'summary':{},'canonical_findings':[finding(evidence,target)]}
    output=build_annotated_docx(data,review)
    assert ranges(output) == [target]
    assert review['canonical_findings'][0]['placement_status'] == 'verified_source_location'


def test_repeated_quote_uses_matching_section_and_full_evidence():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Research Design',2)
    doc.add_paragraph('Participants were selected randomly. The design was described separately.')
    doc.add_heading('Sampling Procedure',2)
    target='Participants were selected randomly. Only available volunteers were approached.'
    doc.add_paragraph(target)
    data=source(doc);rows=parse_document(data,'chapter.docx')
    evidence=next(r for r in rows if r['text']==target);evidence['paragraph']=3
    output=build_annotated_docx(data,{'summary':{},'canonical_findings':[finding(evidence,'Only available volunteers were approached.')]})
    assert ranges(output) == ['Only available volunteers were approached.']


def test_ambiguous_text_without_a_valid_ordinal_is_not_guessed():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    repeated='The sampling procedure selected only volunteers.'
    doc.add_paragraph(repeated);doc.add_paragraph(repeated)
    evidence={'paragraph':999,'text':repeated,'chapter_number':3,'heading':'Sampling Procedure','document_role':'current'}
    source_map,_=_source_locator_map(doc)
    assert resolve_verified_location(finding(evidence,repeated),source_map) == 0


def test_unmatched_quote_stays_in_report_without_a_wrong_body_comment():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    actual='The sample was selected through stratified random sampling.';doc.add_paragraph(actual)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==actual)
    review={'summary':{},'canonical_findings':[finding(evidence,'The researcher recruited only volunteers.')]}
    native=build_annotated_docx(data,review)
    inline=build_inline_annotated_docx(data,review)
    assert not ranges(native)
    assert not any(p.text.startswith('Detailed supervisor comment:') for p in Document(io.BytesIO(inline)).paragraphs)
    assert len(review['canonical_findings']) == 1
    assert review['canonical_findings'][0]['placement_status'] == 'manual_location_required'
    assert review['summary']['manual_comment_location_count'] == 1


@pytest.mark.parametrize('style', ['exact_anchor_grouped', 'one_per_finding'])
def test_two_sentence_findings_do_not_span_the_unrelated_middle_sentence(monkeypatch,style):
    monkeypatch.setenv('VPROF_NATIVE_COMMENT_STYLE',style)
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    first='The sample consisted only of volunteers.';middle='The instrument contained demographic questions.';last='The sampling procedure is described as random.'
    text=first+' '+middle+' '+last;doc.add_paragraph(text)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==text)
    one=finding(evidence,first,1);two=finding(evidence,last,2)
    two['item']='Clarify the random-selection claim';two['required_action']='Explain how random selection was implemented.'
    output=build_annotated_docx(data,{'summary':{},'canonical_findings':[one,two]})
    assert sorted(ranges(output)) == sorted([first,last])


def test_wrong_chapter_number_cannot_override_exact_content():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    text='The sampling procedure selected only volunteers.';doc.add_paragraph(text)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==text)
    evidence['chapter_number']=4
    review={'summary':{},'canonical_findings':[finding(evidence,text)]}
    assert not ranges(build_annotated_docx(data,review))


@pytest.mark.parametrize('style', ['exact_anchor_grouped', 'one_per_finding'])
def test_table_finding_attaches_to_quoted_cell_in_verified_row(monkeypatch,style):
    monkeypatch.setenv('VPROF_NATIVE_COMMENT_STYLE',style)
    doc=Document();doc.add_heading('CHAPTER FOUR',1);doc.add_heading('Regression Results',2)
    doc.add_paragraph('Table 4.1 Coefficients')
    table=doc.add_table(rows=3,cols=3)
    for cells,values in zip(table.rows,[['Variable','B','p'],['Age','.2','.04'],['Income','.5','.01']]):
        for cell,value in zip(cells.cells,values):cell.text=value
    data=source(doc);evidence=next(r for r in parse_document(data,'results.docx') if r.get('table_cells',[])[:1]==['Age'])
    row=finding(evidence,'.04');row['category']='statistical_accuracy';row['item']='Verify the probability for age';row['required_action']='Verify the p-value for age against the model output.'
    output=build_annotated_docx(data,{'summary':{},'canonical_findings':[row]})
    assert ranges(output) == ['.04']


def test_inline_note_follows_the_verified_paragraph_after_ordinal_drift():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    doc.add_paragraph('An unrelated description appears here.')
    text='The sampling procedure selected only volunteers.';doc.add_paragraph(text)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==text);evidence['paragraph']-=1
    output=build_inline_annotated_docx(data,{'summary':{},'canonical_findings':[finding(evidence,text)]})
    paragraphs=Document(io.BytesIO(output)).paragraphs
    i=next(i for i,p in enumerate(paragraphs) if p.text.startswith('Detailed supervisor comment:'))
    assert paragraphs[i-1].text.startswith(text)


@pytest.mark.parametrize('spec_aligned', ['true', 'false'])
def test_unverified_location_has_full_correction_in_report_even_when_details_disabled(monkeypatch,spec_aligned):
    from app.report_exporter import build_docx_report
    monkeypatch.setenv('VPROF_SPEC_ALIGNED_SUPERVISORY_REPORT',spec_aligned)
    monkeypatch.setenv('VPROF_REPORT_INCLUDE_DETAILED_FINDINGS','false')
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    actual='The sample was selected through stratified random sampling.';doc.add_paragraph(actual)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==actual)
    review={'summary':{'filename':'chapter.docx','academic_level':'PhD'},'canonical_findings':[finding(evidence,'The researcher recruited only volunteers.')]}
    build_annotated_docx(data,review)
    report=Document(io.BytesIO(build_docx_report(review)))
    text=' '.join([p.text for p in report.paragraphs]+[cell.text for t in report.tables for r in t.rows for cell in r.cells])
    assert 'manual location checking' in text
    assert 'Correct the sampling description using the actual procedure.' in text


def test_inline_table_note_follows_the_quoted_cell_in_verified_row():
    doc=Document();doc.add_heading('CHAPTER FOUR: RESULTS',1);doc.add_heading('Regression Results',2)
    doc.add_paragraph('Table 4.1 Coefficients');table=doc.add_table(rows=3,cols=3)
    for cells,values in zip(table.rows,[['Variable','B','p'],['Age','.2','.04'],['Income','.5','.01']]):
        for cell,value in zip(cells.cells,values):cell.text=value
    data=source(doc);evidence=next(r for r in parse_document(data,'results.docx') if r.get('table_cells',[])[:1]==['Age'])
    row=finding(evidence,'.04');row['category']='statistical_accuracy'
    output=Document(io.BytesIO(build_inline_annotated_docx(data,{'summary':{},'canonical_findings':[row]})))
    cell=output.tables[0].rows[1].cells[2]
    assert cell.paragraphs[0].text.startswith('.04')
    assert cell.paragraphs[1].text.startswith('Detailed supervisor comment:')
    assert all('Detailed supervisor comment:' not in c.text for c in output.tables[0].rows[2].cells)


def test_unverified_offsets_cannot_override_the_source_quote():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    first='The instrument contained demographic questions.';last='The sample consisted only of volunteers.'
    text=first+' '+last;doc.add_paragraph(text)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==text)
    row=finding(evidence,last);row.update({'exact_anchor_start':0,'exact_anchor_end':len(first)})
    assert ranges(build_annotated_docx(data,{'summary':{},'canonical_findings':[row]})) == [last]


def test_duplicate_passages_are_not_disambiguated_by_a_valid_but_stale_ordinal():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    repeated='The sampling procedure selected only volunteers.'
    doc.add_paragraph(repeated);doc.add_paragraph(repeated)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==repeated)
    review={'summary':{},'canonical_findings':[finding(evidence,repeated)]}
    assert not ranges(build_annotated_docx(data,review))
    assert review['summary']['manual_comment_location_count'] == 1


def test_repeated_phrase_within_one_paragraph_requires_a_more_precise_quote():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Sampling Procedure',2)
    text='Participants volunteered for the pilot. Participants volunteered for the final study.';doc.add_paragraph(text)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==text)
    review={'summary':{},'canonical_findings':[finding(evidence,'Participants volunteered')]}
    assert not ranges(build_annotated_docx(data,review))


def test_exact_content_in_the_wrong_section_does_not_override_section_evidence():
    doc=Document();doc.add_heading('CHAPTER THREE',1);doc.add_heading('Research Design',2)
    text='The sampling procedure selected only volunteers.';doc.add_paragraph(text)
    data=source(doc);evidence=next(r for r in parse_document(data,'chapter.docx') if r['text']==text)
    evidence['heading']='Sampling Procedure'
    review={'summary':{},'canonical_findings':[finding(evidence,text)]}
    assert not ranges(build_annotated_docx(data,review))
