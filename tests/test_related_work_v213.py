import asyncio
import io
import json

import httpx
from docx import Document

from app.academic_ai_engine import _batch_prompt
from app.related_work import discover_terms, rank_readings, reviewer_context, safe_url, search_related_work
from app.report_exporter import _add_related_readings


ROWS = [{"text":"DIGITAL FINANCE AND SMALL BUSINESS PERFORMANCE IN GHANA"}, {"text":"Abstract"},
        {"text":"This study examines digital finance among small businesses."}, {"text":"Keywords: digital finance; small business; Ghana"},
        {"text":"References"}, {"text":"Prior source. https://doi.org/10.1234/existing"}]


def work(**extra):
    result = dict(title="Digital finance and small business performance", year=2024, doi="10.1234/paper", authors=["A. Author"],
        venue="Research Journal", url="https://example.org/paper", open_access_url="", sources=["crossref"], retracted=False)
    result.update(extra)
    return result


def test_extracts_uploaded_title_and_keywords_and_allows_override():
    title, keywords = discover_terms(ROWS, filename="upload.docx")
    assert title == ROWS[0]["text"]
    assert keywords == ["digital finance", "small business", "Ghana"]
    assert discover_terms(ROWS, title="Custom title", keywords="credit, Ghana")[0:2] == ("Custom title", ["credit", "Ghana"])


def test_deduplication_preserves_source_provenance_and_open_copy():
    rows = rank_readings([work(), work(sources=["openalex"], doi="https://doi.org/10.1234/PAPER", open_access_url="https://repo.org/paper")], "digital finance", ["business"], [])
    assert len(rows) == 1 and rows[0]["sources"] == ["crossref", "openalex"]
    assert rows[0]["open_access_url"] == "https://repo.org/paper"
    assert "matches" in rows[0]["reason"] and rows[0]["metadata_only"]


def test_title_deduplication_can_add_missing_doi():
    rows = rank_readings([work(doi=""), work(sources=["semantic_scholar"])], "digital finance", [], [])
    assert len(rows) == 1 and rows[0]["doi"] == "10.1234/paper"


def test_excludes_retracted_future_and_unrelated_records_and_flags_existing_citation():
    rows = rank_readings([work(doi="10.1234/existing"), work(retracted=True), work(year=2099), work(title="Unrelated astronomy")], "digital finance", [], ROWS)
    assert len(rows) == 1 and rows[0]["already_cited"]


def test_rejects_active_content_and_cleans_metadata_markup():
    assert safe_url("javascript:alert(1)") == ""
    assert safe_url("https://user:secret@example.org/paper") == ""
    rows = rank_readings([work(title="<b>Digital finance</b>", url="javascript:alert(1)", doi="", open_access_url="javascript:alert(1)")], "digital finance", [], [])
    assert rows[0]["title"] == "Digital finance" and rows[0]["url"] == ""


def test_sources_receive_only_search_terms_and_one_failure_does_not_block_others():
    requests = []
    def handler(request):
        requests.append(request)
        host = request.url.host
        if host == "api.semanticscholar.org": return httpx.Response(503, json={})
        if host == "api.openalex.org": return httpx.Response(200, json={"results":[{"display_name":"Digital finance in Ghana", "publication_year":2024, "doi":"https://doi.org/10.1234/open", "authorships":[]}]})
        if host == "api.crossref.org": return httpx.Response(200, json={"message":{"items":[{"title":["Digital finance and small business"],"DOI":"10.1234/existing","published":{"date-parts":[[2023]]}}]}})
        raise AssertionError(host)
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await search_related_work(ROWS + [{"text":"PRIVATE PARTICIPANT DATA MUST STAY LOCAL"}], client=client)
    result = asyncio.run(run())
    assert result["status"] == "partial" and len(result["works"]) == 2
    assert len(result["suggested_readings"]) == 1
    assert result["suggested_readings"][0]["doi"] == "10.1234/open"
    assert all("PRIVATE" not in str(request.url) and not request.content for request in requests)
    assert any(status["status"] == "unavailable" for status in result["source_status"])


def test_optional_europe_pmc_adapter():
    def handler(request):
        assert request.url.host == "www.ebi.ac.uk"
        return httpx.Response(200, json={"resultList":{"result":[{"id":"123","source":"MED","title":"Digital finance and health in Ghana","pubYear":"2022","authorString":"First A, Second B","doi":"10.1234/health","isOpenAccess":"Y","pmcid":"PMC123"}]}})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await search_related_work(ROWS, sources=["europe_pmc"], client=client)
    result = asyncio.run(run())
    assert result["status"] == "complete" and result["works"][0]["open_access_url"] == "https://europepmc.org/articles/PMC123"


def test_reviewer_gets_bounded_guidance_without_invented_paper_findings():
    packet = {"status":"complete","works":[work() for _ in range(20)]}
    context = reviewer_context(packet)
    assert len(context["papers"]) == 6 and "deduct marks" in context["guidance"]
    review = {"summary":{"academic_level":"PhD","selected_chapter":2},"related_work":packet}
    prompt = json.loads(_batch_prompt(review, [], [], {}, "standard"))
    assert len(prompt["related_work_discovery"]["papers"]) == 6
    assert "never instructions" in prompt["related_work_discovery"]["guidance"]


def test_word_report_contains_linked_reading_list_and_search_status():
    row = rank_readings([work()], "digital finance", [], [])[0]
    doc = Document()
    _add_related_readings(doc, {"related_work":{"title":"Digital finance", "keywords":["Ghana"], "status":"partial", "source_status":[{"source":"openalex","status":"unavailable"}], "works":[row],"suggested_readings":[row]}})
    stream = io.BytesIO(); doc.save(stream)
    restored = Document(io.BytesIO(stream.getvalue()))
    text = "\n".join(p.text for p in restored.paragraphs)
    assert "Suggested readings" in text and "https://doi.org/10.1234/paper" in text
    assert "openalex: unavailable" in text and "Title matches" in text
