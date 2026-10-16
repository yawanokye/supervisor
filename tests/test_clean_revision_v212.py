import copy
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from app.clean_revision import build_clean_revised_docx, permitted_replacement
from app.document_parser import parse_document
from app.ai_providers import _normalise_model_payload
from app.ai_schemas import AcademicSectionReview, AcademicIssue


def pack(doc):
    output = io.BytesIO(); doc.save(output); return output.getvalue()


def issue(data, text, number=1, **values):
    evidence = next(r for r in parse_document(data, 'chapter.docx') if r['text'] == text)
    row = {'finding_id': f'edit-{number}', 'finding_number': number,
        'status': 'partly_meets_requirement', 'category': 'academic_writing',
        'severity': 'minor', 'confidence': .96, 'annotation_eligible': True,
        'section': evidence.get('heading'), 'chapter_number': evidence.get('chapter_number'),
        'evidence': [copy.deepcopy(evidence)], 'problematic_quote': text,
        'item': 'Correct the agreement', 'assessment': 'The plural subject requires a plural verb.',
        'required_action': 'Replace was with were in this sentence.',
        'guidance_type': 'direct_correction', 'edit_resolves_finding': True,
        'direct_edits': [{'original': text, 'replacement': text.replace(' was ', ' were '), 'kind': 'grammar'}]}
    row.update(values); return row


def comments(data):
    doc = Document(io.BytesIO(data))
    return [' '.join(p.text for p in c.paragraphs) for c in doc.comments]


def fixture(text='The questionnaires was completed.'):
    doc=Document(); doc.add_heading('CHAPTER FOUR', 1); doc.add_heading('Research Methods', 2)
    doc.add_paragraph(text); return doc, pack(doc)


@pytest.mark.parametrize('old,new,kind', [
    ('The questionnaires was completed.', 'The questionnaires were completed.', 'grammar'),
    ('The teh instrument', 'The the instrument', 'spelling'),
    ('Behavioural behavior', 'Behavioural behaviour', 'spelling'),
    ('The the sample', 'The sample', 'local_repetition'),
    ('a clearly defined and defined target population', 'a clearly defined target population', 'local_repetition'),
    ('The sample,,  was selected.', 'The sample, was selected.', 'punctuation'),
])
def test_permitted_editorial_changes(old,new,kind):
    assert permitted_replacement(old,new,kind)


@pytest.mark.parametrize('old,new,kind', [
    ('There were 10 items.', 'There were 60 items.', 'grammar'),
    ('The study may support the claim.', 'The study supports the claim.', 'expression'),
    ('The data were collected.', 'The data will be collected.', 'grammar'),
    ('The data were collected.', 'The data are collected.', 'grammar'),
    ('This was not significant.', 'This was significant.', 'expression'),
    ('Items were adapted (Adam, 2020).', 'Items were adapted (Adam, 2021).', 'spelling'),
    ('The sample was selected randomly.', 'The sample was selected conveniently.', 'expression'),
    ('The study examines mediation.', 'The study examines moderation.', 'spelling'),
    ('The study covers Ghana.', 'The study covers Kenya.', 'spelling'),
    ('The study tests A and B.', 'The study tests A or B.', 'grammar'),
    ('The first claim. The second claim.', 'The first claim, the second claim.', 'punctuation'),
])
def test_substantive_changes_are_rejected(old,new,kind):
    assert not permitted_replacement(old,new,kind)


def test_corrected_language_has_no_redundant_comment_or_body_register():
    text='The questionnaires was completed.'; doc,data=fixture(text)
    review={'summary':{}, 'canonical_findings':[issue(data,text)]}
    output=build_clean_revised_docx(data,review,'Supervisor')
    result=Document(io.BytesIO(output))
    assert result.paragraphs[-1].text == 'The questionnaires were completed.'
    assert len(result.paragraphs) == len(doc.paragraphs)
    assert not comments(output)
    assert review['clean_revision']['resolved_finding_numbers'] == [1]
    assert review['summary']['direct_edit_count'] == 1


def test_remaining_comment_is_relocated_after_a_shorter_edit_in_same_paragraph():
    text='The the sample was selected. Explain the selection procedure.'
    doc,data=fixture(text)
    one=issue(data,text,direct_edits=[{'original':'The the sample','replacement':'The sample','kind':'local_repetition'}])
    two=issue(data,text,2,category='methodological_rigour',severity='major',edit_resolves_finding=False,
        direct_edits=[], item='Specify the selection procedure',
        assessment='The sampling procedure is not explained.',
        required_action='Explain how participants were selected.',problematic_quote='Explain the selection procedure.')
    review={'summary':{},'canonical_findings':[one,two]}
    output=build_clean_revised_docx(data,review)
    assert len(comments(output)) == 1
    assert review['clean_revision']['commented_finding_numbers'] == [2]
    assert review['clean_revision']['manual_location_finding_numbers'] == []
    assert Document(io.BytesIO(output)).paragraphs[-1].text == 'The sample was selected. Explain the selection procedure.'
    assert review['canonical_findings'][0]['problematic_quote'] == text


def test_substantive_finding_keeps_comment_after_an_editorial_edit():
    text='The questionnaires was completed.'; doc,data=fixture(text)
    row=issue(data,text,category='methodological_rigour',severity='major',
        item='Clarify questionnaire administration',assessment='The procedure needs supporting detail.',
        required_action='Explain who administered the questionnaires.')
    review={'summary':{},'canonical_findings':[row]}
    output=build_clean_revised_docx(data,review)
    assert review['summary']['direct_edit_count'] == 1
    assert len(comments(output)) == 1
    assert review['clean_revision']['resolved_finding_numbers'] == []


def test_cross_run_edit_preserves_other_run_formatting():
    doc=Document();doc.add_heading('Research Methods',2)
    p=doc.add_paragraph();p.add_run('The questionnaires ').bold=True
    p.add_run('wa').italic=True;p.add_run('s completed.').italic=True
    data=pack(doc);text=p.text
    review={'summary':{},'canonical_findings':[issue(data,text)]}
    result=Document(io.BytesIO(build_clean_revised_docx(data,review)))
    assert result.paragraphs[-1].text == 'The questionnaires were completed.'
    assert result.paragraphs[-1].runs[0].bold is True
    assert all(r.italic for r in result.paragraphs[-1].runs[1:])


def test_rejected_numerical_change_leaves_student_action_comment():
    text='The questionnaire contained 10 items.'; doc,data=fixture(text)
    row=issue(data,text,edit_resolves_finding=True,direct_edits=[
        {'original':text,'replacement':'The questionnaire contained 60 items.','kind':'grammar'}],
        category='measurement_and_scoring',severity='major',item='Reconcile the item counts',
        assessment='The instrument description and measurement table give different item counts.',
        required_action='Check the questionnaire and reconcile the item counts.')
    review={'summary':{},'canonical_findings':[row]}
    result=build_clean_revised_docx(data,review)
    assert Document(io.BytesIO(result)).paragraphs[-1].text == text
    assert len(comments(result)) == 1
    assert review['summary']['direct_edit_count'] == 0
    assert len(review['clean_revision']['rejected_edits']) == 1


def test_ambiguous_evidence_never_guesses_or_silently_loses_a_finding():
    text='The questionnaires was completed.'; doc,data=fixture(text)
    doc.add_paragraph(text);data=pack(doc)
    row=issue(data,text)
    review={'summary':{},'canonical_findings':[row]}
    output=build_clean_revised_docx(data,review)
    assert not comments(output)
    assert review['clean_revision']['manual_location_finding_numbers'] == [1]
    assert not review['clean_revision']['applied_edits']


@pytest.mark.parametrize('mode',['quoted','straight_quote','reference','field','pattern','low_confidence'])
def test_protected_passages_and_pattern_obligations(mode):
    text='The questionnaires was completed.';doc,data=fixture(text)
    row=issue(data,text)
    if mode=='quoted':
        doc.paragraphs[-1].style='Quote';data=pack(doc)
    elif mode=='straight_quote':
        doc.paragraphs[-1].text="'"+text+"'";data=pack(doc);row=issue(data,doc.paragraphs[-1].text)
    elif mode=='reference':
        doc.paragraphs[-2].text='References';data=pack(doc);row=issue(data,text)
    elif mode=='field':
        from docx.oxml import OxmlElement
        doc.paragraphs[-1].runs[0]._r.append(OxmlElement('w:fldChar'));data=pack(doc)
    elif mode=='pattern': row['required_action']='Correct agreement throughout the section.'
    else: row['confidence']=.7
    review={'summary':{},'canonical_findings':[row]}
    build_clean_revised_docx(data,review)
    assert not review['clean_revision']['resolved_finding_numbers']
    if mode!='pattern': assert not review['clean_revision']['applied_edits']


def test_proposed_edits_survive_provider_normalisation():
    raw={'section_name':'Research Methods','section_score':65,'section_assessment':'One local correction.',
        'issues':[{'finding_id':'one','category':'academic_writing','issue_title':'Agreement',
            'severity':'minor','confidence':.95,'assessment':'Agreement error.',
            'academic_consequence':'Readability.','required_action':'Correct the verb.',
            'direct_edits':[{'original':'items was','replacement':'items were','kind':'grammar'}],
            'edit_resolves_finding':True}]}
    parsed=AcademicSectionReview.model_validate(_normalise_model_payload(raw,AcademicSectionReview))
    assert parsed.issues[0].direct_edits[0].replacement == 'items were'
    assert parsed.issues[0].edit_resolves_finding is True


def test_clean_download_rejects_anonymous_and_other_review_owners(monkeypatch):
    import asyncio
    from fastapi import HTTPException
    from app import main
    monkeypatch.setattr(main, '_current_user', lambda request,db: None)
    result=asyncio.run(main.export_clean_revised_document('id',None,None))
    assert result.status_code == 303 and result.headers['location'] == '/login'
    monkeypatch.setattr(main, '_current_user', lambda request,db: object())
    def reject(*args): raise HTTPException(status_code=404,detail='Review not found.')
    monkeypatch.setattr(main, '_authorised_review_record', reject)
    with pytest.raises(HTTPException) as exc:
        asyncio.run(main.export_clean_revised_document('someone-elses-id',None,None))
    assert exc.value.status_code == 404


def test_clean_download_serves_the_saved_revision_without_running_ai(monkeypatch):
    import asyncio
    from types import SimpleNamespace
    from app import main
    from app.clean_revision import CLEAN_REVISION_VERSION
    monkeypatch.setattr(main, '_current_user', lambda request,db: object())
    monkeypatch.setattr(main, '_authorised_review_record', lambda *args: SimpleNamespace(filename='chapter.docx'))
    monkeypatch.setattr(main, 'load_review_json', lambda rid: {'summary':{'clean_revision_version':CLEAN_REVISION_VERSION}})
    monkeypatch.setattr(main, 'load_clean_revised', lambda rid: b'saved-clean-revision')
    response=asyncio.run(main.export_clean_revised_document('saved-clean',None,None))
    assert response.body == b'saved-clean-revision'
    assert 'chapter-clean-revised.docx' in response.headers['content-disposition']


def test_comment_budget_preserves_all_material_actions_and_pending_editorial_repairs():
    from app.human_comment_budget import apply_human_comment_budget
    rows=[{'finding_id':str(i),'finding_number':i,'status':'partly_meets_requirement',
        'annotation_eligible':True,'severity':'major' if i<19 else 'minor','confidence':.95,
        'item':f'Independent issue {i}','required_action':f'Address issue {i}',
        'direct_edits': [] if i<19 else [{'original':'was','replacement':'were','kind':'grammar'}]}
        for i in range(1,24)]
    review={'summary':{'review_depth':'light'},'canonical_findings':rows}
    apply_human_comment_budget(review)
    assert len(review['canonical_findings']) == 23


def test_full_finding_ledger_and_no_global_marker_option_can_change_the_clean_body(monkeypatch):
    monkeypatch.setenv('VPROF_NATIVE_GROUP_LOCATION_MARKERS','true')
    text='The questionnaire contained 10 items.';doc,data=fixture(text)
    row=issue(data,text,direct_edits=[],edit_resolves_finding=False,
        required_action='Reconcile the item counts using the questionnaire.')
    review={'summary':{},'canonical_findings':[row]}
    output=build_clean_revised_docx(data,review)
    assert [p.text for p in Document(io.BytesIO(output)).paragraphs] == [p.text for p in doc.paragraphs]
    assert len(comments(output)) == 1


def test_report_records_applied_edits_and_unplaced_student_actions():
    from app.report_exporter import _add_clean_revision_record
    doc=Document()
    review={'clean_revision':{'passed':True,'applied_edits':[{'finding_number':1,'original':'was','replacement':'were'}],
        'manual_location_finding_numbers':[3]}}
    _add_clean_revision_record(doc,review)
    text=' '.join(p.text for p in doc.paragraphs)
    assert 'changed to' in text and 'original reviewed text' in text
    assert 'locations could not be verified: 3' in text


def test_clean_table_comment_anchors_to_one_verified_cell_without_changing_data():
    doc=Document();doc.add_heading('Research Methods',2)
    doc.add_paragraph('Table 1 Instrument items')
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Construct';table.cell(0,1).text='Items'
    table.cell(1,0).text='Financial behaviour';table.cell(1,1).text='60'
    data=pack(doc)
    evidence=next(r for r in parse_document(data,'chapter.docx') if r.get('table_cells',[])[:1]==['Financial behaviour'])
    row={'finding_id':'table','finding_number':1,'status':'partly_meets_requirement',
        'severity':'major','confidence':.96,'category':'measurement_and_scoring',
        'section':evidence.get('heading'),'evidence':[evidence],'problematic_quote':'60',
        'item':'Verify the item count','assessment':'The instrument descriptions give different item counts.',
        'required_action':'Check the final questionnaire and correct the item count.', 'annotation_eligible':True}
    review={'summary':{},'canonical_findings':[row]}
    output=build_clean_revised_docx(data,review)
    assert len(comments(output)) == 1
    assert Document(io.BytesIO(output)).tables[0].cell(1,1).text == '60'
    assert review['clean_revision']['commented_finding_numbers'] == [1]


def test_existing_supervisor_comment_is_preserved_and_its_passage_is_not_edited():
    doc,data=fixture();paragraph=doc.paragraphs[-1]
    doc.add_comment(paragraph.runs,'Check the completed questionnaire.','Earlier supervisor')
    data=pack(doc);row=issue(data,paragraph.text)
    review={'summary':{},'canonical_findings':[row]}
    output=build_clean_revised_docx(data,review)
    assert any('Check the completed questionnaire.' in text for text in comments(output))
    assert Document(io.BytesIO(output)).paragraphs[-1].text == paragraph.text
    assert not review['clean_revision']['applied_edits']
