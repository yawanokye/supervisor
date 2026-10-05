from __future__ import annotations

import asyncio
import copy
import io
import json

import fitz
import pytest
from docx import Document
from pydantic import BaseModel

from app.ai_config import HybridAIConfig
from app.ai_providers import AIProviderError, OpenAIProvider, ProviderResult
from app.ai_schemas import AIUsageRecord
from app.conceptual_framework_review import is_quantitative_study
from app.document_parser import parse_document
from app.guided_review import merge_guided_reviews
from app.human_comment_budget import apply_human_comment_budget
from app.model_compatibility import capabilities, compatible_effort
from app.provider_state import provider_checkpoint_context
from app.results_ledger import run_final_consistency_audit
from app.source_cache import parse_document_cached
from app.statistical_math import recomputed_test_p_mismatch, rounded_ratio_matches
from app.statistical_review import audit_statistical_consistency, audit_table_level_accuracy
from app.supervisory_voice import usable_narrative, comment_variation_audit


class Payload(BaseModel):
    judgement: str


class MemoryCheckpoints:
    def __init__(self): self.values = {}
    def load(self, key, expected_input_hash=''):
        value = self.values.get(key)
        if value and (not expected_input_hash or value[0] == expected_input_hash):
            return copy.deepcopy(value[1])
    def save(self, key, data, input_hash='', **kwargs):
        self.values[key] = (input_hash,copy.deepcopy(data))
    def load_provider_result(self, key, expected_input_hash=''):
        return self.load(key,expected_input_hash)
    def save_provider_result(self, key, result, input_hash='', **kwargs):
        self.save(key,result,input_hash)


def config(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','test-only-key')
    monkeypatch.setenv('AI_STRUCTURED_OUTPUT_RETRIES','0')
    monkeypatch.setenv('OPENAI_BACKGROUND_MODE','false')
    return HybridAIConfig.from_env()


def response(text='{"judgement":"valid"}'):
    return {'id':'resp-test','status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':text}]}],
        'usage':{'input_tokens':10,'output_tokens':12}}


def row(text, number=1, heading='Results', chapter=4):
    return {'paragraph':number,'text':text,'heading':heading,'chapter_number':chapter,'document_role':'current'}


def test_global_model_switches_all_roles_including_retained_overrides(monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL','gpt-6.1-sol')
    monkeypatch.setenv('OPENAI_CHAPTER_MODEL','gpt-5.6-luna')
    monkeypatch.setenv('OPENAI_EXPERT_MODEL','gpt-5.6-terra')
    monkeypatch.setenv('OPENAI_MODEL_MODE','single')
    c = HybridAIConfig.from_env()
    assert c.openai_chapter_model == c.openai_expert_model == c.openai_final_synthesis_model == c.openai_external_domain_model == 'gpt-6.1-sol'
    monkeypatch.setenv('OPENAI_MODEL','future-review-model')
    assert HybridAIConfig.from_env().openai_chapter_model == 'future-review-model'


def test_roles_mode_preserves_intentional_tiers(monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL','gpt-6.1-sol')
    monkeypatch.setenv('OPENAI_MODEL_MODE','roles')
    monkeypatch.setenv('OPENAI_CHAPTER_MODEL','gpt-5.6-luna')
    monkeypatch.delenv('OPENAI_EXPERT_MODEL',raising=False)
    c = HybridAIConfig.from_env()
    assert c.openai_chapter_model == 'gpt-5.6-luna'
    assert c.openai_expert_model == 'gpt-6.1-sol'


def test_effort_matches_model_capability():
    assert compatible_effort('gpt-6.1-sol','none') == 'low'
    assert compatible_effort('gpt-4.1','xhigh') is None
    assert compatible_effort('o3-mini','xhigh') == 'high'
    with pytest.raises(ValueError,match='not a text-review model'): capabilities('text-embedding-3-large')


def test_explicit_parameter_rejection_adapts_without_model_substitution(monkeypatch):
    captured=[]
    async def post(**kwargs):
        captured.append(copy.deepcopy(kwargs['payload']))
        if len(captured) == 1: raise AIProviderError('HTTP 400: text.format json_schema not supported')
        return response(),'request-test'
    monkeypatch.setattr('app.ai_providers._post_json_with_retry',post)
    result=asyncio.run(OpenAIProvider(config(monkeypatch)).complete_json(model='future-review-model',system_prompt='Return JSON.',user_prompt='Review.',schema_model=Payload,purpose='compatibility'))
    assert result.data['judgement'] == 'valid'
    assert [b['model'] for b in captured] == ['future-review-model']*2
    assert 'reasoning' not in captured[0]
    assert captured[1]['text']['format']['type'] == 'json_object'


def test_auth_error_is_not_retried_as_schema_repair(monkeypatch):
    count=0
    async def post(**kwargs):
        nonlocal count
        count+=1
        raise AIProviderError('HTTP 401: invalid key')
    monkeypatch.setattr('app.ai_providers._post_json_with_retry',post)
    monkeypatch.setenv('AI_STRUCTURED_OUTPUT_RETRIES','2')
    c=HybridAIConfig.from_env()
    with pytest.raises(AIProviderError,match='401'):
        asyncio.run(OpenAIProvider(c).complete_json(model='gpt-6.1-sol',system_prompt='JSON',user_prompt='Review',schema_model=Payload,purpose='auth'))
    assert count == 1


def test_chat_models_keep_schema_and_validate_json(monkeypatch):
    captured={}
    async def post(**kwargs):
        captured.update(kwargs)
        return {'choices':[{'message':{'content':'{"judgement":"valid"}'},'finish_reason':'stop'}],'usage':{'prompt_tokens':9,'completion_tokens':5}},'chat-request'
    monkeypatch.setattr('app.ai_providers._post_json_with_retry',post)
    result=asyncio.run(OpenAIProvider(config(monkeypatch)).complete_json(model='gpt-3.5-turbo',system_prompt='Return JSON.',user_prompt='Review',schema_model=Payload,purpose='chat'))
    assert captured['url'].endswith('/chat/completions')
    assert captured['payload']['response_format'] == {'type':'json_object'}
    assert 'reasoning_effort' not in captured['payload']
    assert 'schema' in captured['payload']['messages'][0]['content']
    assert result.usage.input_tokens == 9


def test_background_id_survives_poll_failure_and_resume(monkeypatch):
    c=config(monkeypatch)
    from dataclasses import replace
    c=replace(c,openai_background_mode_enabled=True)
    manager=MemoryCheckpoints(); posts=0; polls=0
    async def post(**kwargs):
        nonlocal posts
        posts+=1
        return {'id':'resp-retained','status':'queued','output':[]},'submitted-request'
    async def poll(**kwargs):
        nonlocal polls
        polls+=1
        assert kwargs['response_payload']['id'] == 'resp-retained'
        assert any(v[1].get('payload',{}).get('id') == 'resp-retained' for v in manager.values.values())
        if polls == 1: raise AIProviderError('temporary polling timeout')
        return response(),'polled-request'
    monkeypatch.setattr('app.ai_providers._post_json_with_retry',post)
    monkeypatch.setattr('app.ai_providers._poll_openai_background_response',poll)
    async def run():
        with provider_checkpoint_context(manager,'stable-stage'):
            return await OpenAIProvider(c).complete_json(model='gpt-6.1-sol',system_prompt='JSON',user_prompt='Review',schema_model=Payload,purpose='resume',reasoning_effort='high')
    with pytest.raises(AIProviderError,match='timeout'): asyncio.run(run())
    assert asyncio.run(run()).data['judgement'] == 'valid'
    assert posts == 1 and polls == 2


def test_parse_cache_reuses_extraction_and_returns_independent_rows(monkeypatch):
    count=0
    def parse(*args):
        nonlocal count
        count+=1
        return [row('Readable research evidence')]
    monkeypatch.setattr('app.source_cache.parse_document',parse)
    manager=MemoryCheckpoints()
    first=asyncio.run(parse_document_cached(b'doc','thesis.docx',manager)); first[0]['text']='mutated'
    assert asyncio.run(parse_document_cached(b'doc','thesis.docx',manager))[0]['text'] == 'Readable research evidence'
    assert count == 1
    asyncio.run(parse_document_cached(b'changed','thesis.docx',manager))
    assert count == 2


def test_negative_adjusted_r_squared_is_not_impossible():
    warnings=audit_statistical_consistency([row('Adjusted R squared = -0.10.')])
    assert 'invalid_r_squared' not in {w['kind'] for w in warnings}


def test_declared_alpha_and_separate_decisions():
    evidence=[row('The significance level was .01.',1,'Methods',3),row('The association was not significant, p = .04; the second was significant, p = .001.',2)]
    assert 'p_value_interpretation_mismatch' not in {w['kind'] for w in audit_statistical_consistency(evidence)}
    evidence[1]['text']='The association was significant, p = .04.'
    assert 'p_value_interpretation_mismatch' in {w['kind'] for w in audit_statistical_consistency(evidence)}


def test_signed_statistic_mismatch_and_rounding():
    assert any(w['kind']=='coefficient_se_t_mismatch' for w in audit_statistical_consistency([row('B = -.50, SE = .10, t = 5.0.')]))
    assert rounded_ratio_matches('.005','.001','4.0')
    assert not rounded_ratio_matches('-.50','.10','5.0')
    assert rounded_ratio_matches('-.50','.10','5.0',absolute_statistic=True)


@pytest.mark.parametrize('text,wrong',[
    ('t(28) = 2.048, p = .050',False),('t(28) = 2.048, p < .001',True),
    ('F(2, 97) = 3.09, p = .05',False),('chi-square(1) = 10.83, p > .05',True),
    ('One-tailed t(28) = 2.048, p = .025',False),('Bootstrap t(28) = 2.048, p = .02',False)])
def test_p_recomputation_has_explicit_assumptions(text,wrong):
    assert bool(recomputed_test_p_mismatch(text)) == wrong


def test_logistic_reporting_does_not_require_ols_t_statistics():
    title='Table 4.1 Logistic regression'
    evidence=[{**row(text,i),'source_kind':'table_row','table_index':1,'table_row':i,'table_title':title,'table_number':'4.1','table_cells':text.split(' | ')}
        for i,text in enumerate(['Predictor | B | SE | z | p','Age | .20 | .10 | 2.00 | .045'],1)]
    assert 'regression_table_incomplete' not in {w['kind'] for w in audit_table_level_accuracy(evidence)}


def test_quantitative_detection_uses_words_not_substrings():
    assert not is_quantitative_study([row('A qualitative exploration of semantics in schools.')],research_approach='qualitative')


def test_blank_docx_cells_remain_in_their_columns():
    doc=Document();doc.add_heading('CHAPTER FOUR',0);doc.add_paragraph('Table 4.1 Coefficients')
    table=doc.add_table(rows=2, cols=4)
    for cell,value in zip(table.rows[0].cells,['Variable','B','SE','p']):cell.text=value
    for cell,value in zip(table.rows[1].cells,['Age','.2','','.04']):cell.text=value
    buffer=io.BytesIO();doc.save(buffer)
    rows=parse_document(buffer.getvalue(),'evidence.docx')
    result=next(r for r in rows if r.get('table_cells',[])[:1] == ['Age'])
    assert result['table_cells'] == ['Age','.2','','.04']


def test_pdf_geometry_recovers_table_columns_and_renders_framework():
    pdf=fitz.open();page=pdf.new_page()
    page.insert_text((50,40),'CHAPTER FOUR'); page.insert_text((50,65),'Table 4.1 Regression estimates')
    for y in (90,120,150):page.draw_line((50,y),(350,y))
    for x in (50,150,250,350):page.draw_line((x,90),(x,150))
    for pos,text in [((60,110),'Variable'),((160,110),'B'),((260,110),'p'),((60,140),'Age'),((160,140),'.20'),((260,140),'.04')]:page.insert_text(pos,text)
    page=pdf.new_page();page.insert_text((50,40),'Conceptual Framework');page.insert_text((50,80),'Predictor');page.insert_text((250,80),'Outcome');page.draw_line((120,77),(245,77))
    rows=parse_document(pdf.tobytes(),'geometry.pdf');pdf.close()
    assert any(r.get('table_cells') == ['Age','.20','.04'] for r in rows)
    assert any(r.get('drawing_image_data_urls') for r in rows)


def test_budget_does_not_refill_duplicate_comments_and_keeps_full_ledger():
    finding={'status':'partly_meets_requirement','severity':'moderate','confidence':.9,'category':'academic_writing','item':'Ambiguous construct name','required_action':'Define the construct consistently.'}
    review=apply_human_comment_budget({'summary':{'review_depth':'standard'},'canonical_findings':[{**finding,'finding_id':str(i),'finding_number':i+1} for i in range(3)]})
    assert len(review['canonical_findings']) == 1
    assert len(review['internal_issue_ledger']['findings']) == 3


def test_narrative_cannot_drop_second_action_or_invent_numbers():
    finding={'item':'Objectives use different terminology','required_action':'Align the objectives and questions. Use the same terminology throughout.',
        'supervisory_comment':'Align the objectives and questions.'}
    assert usable_narrative(finding) == ''
    finding['supervisory_comment']='Align the objectives and questions, and use the same terminology throughout.'
    assert usable_narrative(finding)
    finding['supervisory_comment'] += ' Replace p = .05 with .01.'
    assert usable_narrative(finding) == ''


def test_variation_audit_detects_monotonic_openings():
    rows=[{'item':'Different variable '+str(i),'assessment':'This section needs a clear explanation of variable '+str(i)+'.','required_action':'Explain the variable '+str(i)+'.'} for i in range(5)]
    assert comment_variation_audit(rows)['repeated_openings']


def test_chapter_five_only_is_not_complete_or_finally_audited():
    merged=merge_guided_reviews([{'summary':{'selected_chapter':5},'academic_findings':[]}],expected_chapters=[1,2,3,4,5])
    assert merged['summary']['document_label'] != 'Complete thesis'
    assert merged['summary']['guided_unreviewed_chapters'] == [1,2,3,4]
    assert merged['summary']['guided_final_consistency_audit'] is False


def test_final_audit_requires_real_coverage_and_reuses_result(monkeypatch):
    evidence=[row('The study examined age and wellbeing.',1,'Research objectives',1),row('Age was associated with wellbeing.',2,'Conclusions',5)]
    calls=0
    async def complete(self,**kwargs):
        nonlocal calls
        calls+=1
        return ProviderResult(data={'reviews':[{'section_key':'FINAL-CONSISTENCY','section_assessment':'The reported association is consistent with the objective and bounded conclusion.','issues':[],'assessed_paragraph_ids':['P1','P2']}]},usage=AIUsageRecord(provider='openai',model='gpt-6.1-sol',purpose='final'))
    monkeypatch.setattr('app.model_router.CostAwareAIProvider.complete_json',complete)
    manager=MemoryCheckpoints();c=config(monkeypatch)
    for _ in range(2):
        result=asyncio.run(run_final_consistency_audit({'summary':{'guided_chapters_completed':[1,5]}},evidence,manager,c))
        assert result['summary']['guided_final_consistency_audit'] is True
        assert result['results_ledger']['raw_data_reproduced'] is False
    assert calls == 1


def test_empty_final_audit_is_never_labelled_completed(monkeypatch):
    async def empty(self,**kwargs):
        return ProviderResult(data={'reviews':[]},usage=AIUsageRecord(provider='openai',model='test',purpose='final'))
    monkeypatch.setattr('app.model_router.CostAwareAIProvider.complete_json',empty)
    review={'summary':{'guided_final_consistency_audit':False}}
    with pytest.raises(ValueError,match='no valid substantive'):
        asyncio.run(run_final_consistency_audit(review,[row('Test conclusion.',1,'Conclusions')],MemoryCheckpoints(),config(monkeypatch)))
    assert review['summary']['guided_final_consistency_audit'] is False


def test_verified_comment_survives_native_word_export():
    from app.annotated_exporter import build_annotated_docx
    from app.natural_supervisor_comment import natural_supervisor_comment
    import zipfile
    doc=Document();doc.add_heading('CHAPTER ONE',1);doc.add_heading('Research Objectives',2)
    doc.add_paragraph('The objective examines access, while the question concerns participation.')
    buffer=io.BytesIO();doc.save(buffer);source=buffer.getvalue()
    evidence=next(r for r in parse_document(source,'sample.docx') if r['text'].startswith('The objective'))
    finding={'finding_id':'VOICE-1','finding_number':1,'item':'Objective and question use different constructs',
        'category':'objectives_questions_hypotheses','status':'partly_meets_requirement','severity':'major',
        'confidence':.9,'section':'Research Objectives','section_reference':'Research Objectives',
        'assessment':'The objective examines access, but the question concerns participation.',
        'required_action':'Align the objective and question around one construct. Use the same terminology throughout.',
        'supervisory_comment':'Access and participation carry different meanings here. Align the objective and question around one construct, and use the same terminology throughout.',
        'problematic_quote':evidence['text'],'evidence':[evidence],'annotation_eligible':True}
    from app.supervisory_voice import narrative_basis
    finding['_supervisory_comment_basis']=narrative_basis(finding)
    expected=natural_supervisor_comment(finding)
    output=build_annotated_docx(source,{'summary':{'academic_level':'Bachelors'},'canonical_findings':[finding],'academic_findings':[finding]})
    with zipfile.ZipFile(io.BytesIO(output)) as z:
        comments=z.read('word/comments.xml').decode()
    assert expected in comments
    assert 'same terminology throughout' in comments


def test_prose_basis_rejects_a_comment_after_a_factual_correction():
    from app.supervisory_voice import narrative_basis
    finding={'item':'Sample not justified','required_action':'Justify the sample size.',
        'supervisory_comment':'Justify the sample size for the stated analysis.'}
    finding['_supervisory_comment_basis']=narrative_basis(finding)
    assert usable_narrative(finding)
    finding['assessment']='The sample size justification is adequate.'
    assert not usable_narrative(finding)


def test_prose_repair_changes_only_voice_and_preserves_evidence(monkeypatch):
    from app.comment_prose_quality import polish_selected_comments
    finding={'finding_id':'voice-test','item':'Different construct labels','severity':'moderate',
        'status':'partly_meets_requirement','required_action':'Align the objective and question. Use the same terminology throughout.',
        'assessment':'The objective and question name different constructs.', 'evidence':[{'paragraph':4}]}
    original=copy.deepcopy(finding)
    async def complete(self,**kwargs):
        return ProviderResult(data={'comments':[{'finding_id':'voice-test','supervisory_comment':'The construct shifts between the two statements. Align the objective and question, and use the same terminology throughout.'}]},
            usage=AIUsageRecord(provider='openai',model='test',purpose='voice',input_tokens=10,output_tokens=20))
    monkeypatch.setattr('app.model_router.CostAwareAIProvider.complete_json',complete)
    result=asyncio.run(polish_selected_comments({'canonical_findings':[finding]},MemoryCheckpoints(),config(monkeypatch)))
    for key in ('assessment','required_action','severity','evidence'): assert finding[key] == original[key]
    assert result['comment_quality_audit']['accepted_revisions'] == 1
    assert result['ai_review']['api_call_count'] == 1


def test_explicit_stop_cancels_only_pending_background_responses(monkeypatch):
    from types import SimpleNamespace
    from app.provider_cancellation import cancel_retained_background_responses
    from contextlib import contextmanager
    records=[SimpleNamespace(job_id='job-chapter-1',stage_key='provider-inflight-pending'),
        SimpleNamespace(job_id='job-chapter-1',stage_key='provider-inflight-completed')]
    class Query:
        def filter(self,*args): return self
        def all(self): return records
    @contextmanager
    def session():
        yield SimpleNamespace(query=lambda *args:Query())
    manager=MemoryCheckpoints()
    manager.save(records[0].stage_key,{'payload':{'id':'resp-pending','status':'queued'}})
    manager.save(records[1].stage_key,{'payload':{'id':'resp-done','status':'completed'}})
    urls=[]
    async def post(**kwargs):
        urls.append(kwargs['url'])
        return {'id':'resp-pending','status':'cancelled'},'cancel-request'
    monkeypatch.setattr('app.provider_cancellation.SessionLocal',session)
    monkeypatch.setattr('app.provider_cancellation.CheckpointManager',lambda *args:manager)
    monkeypatch.setattr('app.provider_cancellation._post_json_with_retry',post)
    assert asyncio.run(cancel_retained_background_responses('job',config(monkeypatch))) == {'cancelled':1,'failed':0}
    assert urls[0].endswith('/responses/resp-pending/cancel')
    assert len(urls) == 1
    assert manager.load(records[1].stage_key)['payload']['status'] == 'completed'


def test_operator_prices_override_bundled_model_estimates(monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL_PRICES_JSON',json.dumps({'gpt-6.1-sol':{'input':3,'cached_input':.2,'output':12}}))
    assert HybridAIConfig.from_env().openai_prices_for_model('gpt-6.1-sol') == (3,.2,12)


def test_negative_growth_and_non_reliability_alpha_are_not_impossible():
    warnings=audit_statistical_consistency([row('Growth decreased by -20%. The model intercept alpha = -.5.')])
    assert not {'invalid_percentage','negative_reliability_requires_investigation'} & {w['kind'] for w in warnings}


def test_long_shortlist_prose_repair_is_split_into_bounded_batches(monkeypatch):
    from app.comment_prose_quality import polish_selected_comments
    findings=[{'finding_id':f'issue-{n}','item':f'Variable {n} needs a definition','status':'partly_meets_requirement',
        'assessment':f'Variable {n} is named without a definition.','required_action':f'Define variable {n} consistently.'} for n in range(52)]
    sizes=[]
    async def complete(self,**kwargs):
        packet=json.loads(kwargs['user_prompt']);sizes.append(len(packet['findings']))
        return ProviderResult(data={'comments':[{'finding_id':r['finding_id'],'supervisory_comment':r['required_action']} for r in packet['findings']]},
            usage=AIUsageRecord(provider='openai',model='test',purpose='prose',request_id=f'batch-{len(sizes)}'))
    monkeypatch.setattr('app.model_router.CostAwareAIProvider.complete_json',complete)
    result=asyncio.run(polish_selected_comments({'canonical_findings':findings},MemoryCheckpoints(),config(monkeypatch)))
    assert sizes == [24,24,4]
    assert result['comment_quality_audit']['accepted_revisions'] == 52
    assert result['ai_review']['api_call_count'] == 3


def test_unavailable_retained_response_can_be_resubmitted_on_resume(monkeypatch):
    from dataclasses import replace
    c=replace(config(monkeypatch),openai_background_mode_enabled=True)
    manager=MemoryCheckpoints();posts=0;polls=0
    async def post(**kwargs):
        nonlocal posts
        posts+=1
        return {'id':f'resp-{posts}','status':'queued'},'submitted-request'
    async def poll(**kwargs):
        nonlocal polls
        polls+=1
        if polls == 1: raise AIProviderError('HTTP 404: response no longer available')
        return response(),'new-response'
    monkeypatch.setattr('app.ai_providers._post_json_with_retry',post)
    monkeypatch.setattr('app.ai_providers._poll_openai_background_response',poll)
    async def run():
        with provider_checkpoint_context(manager,'stable-stage'):
            return await OpenAIProvider(c).complete_json(model='gpt-6.1-sol',system_prompt='JSON',user_prompt='Review',schema_model=Payload,purpose='expired-recovery',reasoning_effort='high')
    with pytest.raises(AIProviderError,match='404'): asyncio.run(run())
    assert asyncio.run(run()).data['judgement'] == 'valid'
    assert posts == 2
